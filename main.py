import re
import customtkinter as ctk
from tkinter import filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES
import csv, openpyxl
from email_validator import validate_email, EmailNotValidError

class EmailExtractor:
    def __init__(self):
        self.app = TkinterDnD.Tk()
        self.app.title("Email Extractor")
        self.app.geometry("600x700")
        
        ctk.CTkLabel(self.app, text="📧 Email Extractor", 
                    font=("Helvetica", 24, "bold")).pack(pady=20)
        
        ctk.CTkButton(self.app, text="📄 Estrai da File(PDF, CSV, EXCEL E ALTRI)", 
                     command=self.estrai_file).pack(pady=10, padx=50, fill="x")

        
        ctk.CTkButton(self.app, text="📁 Estrai da Cartella", 
                     command=self.estrai_cartella).pack(pady=10, padx=50, fill="x")
        
        ctk.CTkButton(self.app, text="📑 Estrai da URL", 
                     command=self.estrai_url).pack(pady=10, padx=50, fill="x")

        ctk.CTkLabel(self.app, text="-- Oppure usa il Drag and Drop --", 
                    font=("Arial", 16)).pack()
        
        self.textbox = ctk.CTkTextbox(self.app, width=500, height=200, font=("Helvetica", 18))
        #self.textbox.configure(state="disabled")
        self.textbox.pack(pady=20)

        frame_opzioni = ctk.CTkFrame(self.app)
        frame_opzioni.pack(pady=10, padx=30, fill="x")
        
        self.check_duplicati = ctk.CTkCheckBox(frame_opzioni, 
                                               text="Elimina duplicati", command=self.elimina_duplicati)
        self.check_duplicati.pack(side="left", padx=15, pady=10)
        self.check_duplicati.select()  # attivo di default
        


        frame_export = ctk.CTkFrame(self.app)
        frame_export.pack(pady=15, padx=30, fill="x")

        # Frame interno per centrare i widget
        frame_centrato = ctk.CTkFrame(frame_export, fg_color="transparent")
        frame_centrato.pack(anchor="center")

        # Pulsante Salva
        ctk.CTkButton(frame_centrato, text="💾 Salva in", 
                     command=self.salva).pack(side="left", padx=15, pady=10)

        # Combobox formato
        self.formato = ctk.CTkComboBox(frame_centrato, 
                                       values=["CSV", "TXT", "JSON"],
                                       width=100)
        self.formato.pack(side="left", padx=5, pady=10)

        # Pulsante salva
       
        
        self.label_count = ctk.CTkLabel(self.app, text="Email trovate: 0")
        self.label_count.pack()
        
        self.email_trovate = []
        self.email_uniche = []
        self.email_iniziali = []

        self.app.drop_target_register(DND_FILES)
        self.app.dnd_bind('<<Drop>>', self.file_rilasciato)
    
    def trova_email(self, testo):
        """Estrae email dal testo"""
        pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        email = list(re.findall(pattern, testo))
        email.sort()
        return email

    def file_rilasciato(self, event):
        percorso = event.data.strip('{}')
        
        print(f"Percorso: {percorso}")
        if re.search("\.[a-zA-Z0-9]+$", percorso, re.IGNORECASE):
            self.estrai_file( percorso, dialog=False)

        else:
            self.estrai_cartella(percorso, dialog=False)
        
    
    def mostra(self, email):
        """Mostra i risultati"""
        self.textbox.delete("1.0", "end")
        checkbox = self.check_duplicati.get()
        fore = len(self.email_iniziali) == len(email)
        if email:
            if checkbox == 1:
                if fore == True:
                    email = list(set(self.email_iniziali))
                elif fore == False:
                    pass
            elif checkbox == 0:
                if fore == True:
                    email = self.email_iniziali
                elif fore == False:
                    email = self.email_iniziali
                
            self.textbox.insert("1.0", f"Trovate {len(email)} email:\n\n")

            
            for e in email:
                self.textbox.insert("end", f"{e}\n")
            
            self.label_count.configure(text=f"Email trovate: {len(email)}")
            self.email_uniche = self.email_trovate

        else:

            self.textbox.insert("1.0", "Nessuna email trovata.")

            self.label_count.configure(text="Email trovate: 0")
            
            self.email_uniche = []

        


    def elimina_duplicati(self):
        
        email = self.email_iniziali.copy()
        
        #if self.check_duplicati.get():
         #   email = list(set(email))
        #else:
         #   email = self.email_iniziali
        
        
#        email.sort()
 #       self.email_uniche = email
        if email != []:
            self.mostra(email)

                


    
    def estrai_file(self, percorso="null", dialog=True):
        if dialog == True:
            file = filedialog.askopenfilename(
                filetypes=[
            ("Tutti i file supportati", "*.txt *.csv *.md *.log *.xls *.xlsx *.pdf")
                ]
            )
        else:
            file = percorso

        if file:
            if file.endswith("xls") or file.endswith("xlsx"):
                self.estrai_excel(file)
            elif file.endswith("pdf"):
                self.estrai_pdf(file)
            else:
                with open(file, 'r', encoding='utf-8', errors='ignore') as f:
                    testo = f.read()
                self.email_iniziali = self.trova_email(testo)
                print(self.email_iniziali)
                self.mostra(self.email_iniziali)
    
    def estrai_cartella(self, percorso="null", dialog=True):
        import os
        if dialog == True:
            cartella = filedialog.askdirectory()
        else:
            cartella = percorso
        if cartella:
            tutte = []
            for root, dirs, files in os.walk(cartella):
                for file in files:
                    if file.endswith(('.txt', '.csv', '.md', '.log')):
                        with open(os.path.join(root, file), 'r', 
                                 encoding='utf-8', errors='ignore') as f:
                            tutte.extend(self.trova_email(f.read()))
                    elif file.endswith("pdf"):
                        from pypdf import PdfReader
                        reader = PdfReader(file)
                        testo = ""
                        for pagina in reader.pages:
                            testo += pagina.extract_text()
                        tutte.extend(self.trova_email(testo))
                    elif file.endswith("xls") or file.endswith("xlsx"):
                        import openpyxl
                        wb = openpyxl.load_workbook(file)
                        testo = ""
                        for sheet in wb.sheetnames:
                            ws = wb[sheet]
                            for row in ws.iter_rows(values_only=True):
                                for cell in row:
                                    if cell:
                                        testo += str(cell) + " "
                        tutte.extend(self.trova_email(testo))
                        

                        

                            
            self.email_iniziali = tutte                    
            self.mostra(tutte)
    
    def estrai_pdf(self, path):

        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            testo = ""
            for pagina in reader.pages:
                testo += pagina.extract_text()
            self.mostra(self.trova_email(testo))
        except ImportError:
            messagebox.showerror("Errore", "Installa pypdf: pip install pypdf")


    def salva(self):
        if len(self.email_trovate) == 0:
            messagebox.warning("Attenzione", "Nessuna email da salvare")
        
        formato = self.formato.get()
        
        if formato == "CSV":
            file = filedialog.asksaveasfilename(defaultextension=".csv")
            if file:
                with open(file, 'w', newline='') as f:
                    writer = csv.writer(f)
                    writer.writerow(["Email"])
                    for e in self.email_uniche:
                        writer.writerow([e])
        
        elif formato == "TXT":
            file = filedialog.asksaveasfilename(defaultextension=".txt")
            if file:
                with open(file, 'w') as f:
                    for e in self.email_uniche:
                        f.write(e + "\n")
        
        elif formato == "JSON":
            file = filedialog.asksaveasfilename(defaultextension=".json")
            if file:
                import json
                with open(file, 'w') as f:
                    json.dump(self.email_uniche, f, indent=4)
        
        messagebox.showinfo("OK", "Salvato!")

    def estrai_excel(self, path):

        try:
            import openpyxl
            wb = openpyxl.load_workbook(path)
            testo = ""
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if cell:
                            testo += str(cell) + " "
            self.email_trovate = self.trova_email(testo)
            self.mostra(self.email_trovate)
        except ImportError:
            messagebox.showerror("Errore", "Installa openpyxl: pip install openpyxl")

    def estrai_url(self):
        # Finestra input URL
        dialog = ctk.CTkInputDialog(text="Inserisci URL:", title="Estrazione da URL")
        url = dialog.get_input()
        if url:
            try:
                import requests
                response = requests.get(url, timeout=10)
                self.email_iniziali = self.trova_email(response.text)
                
                self.mostra(self.email_trovate)
            except Exception as e:
                messagebox.showerror("Errore", f"Impossibile scaricare: {e}")
    
    
    def run(self):
        self.app.mainloop()

if __name__ == "__main__":
    app = EmailExtractor()
    app.run()
